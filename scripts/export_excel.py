"""Export literal Excel inputs to a private, reviewable import file. Never edits Excel."""
import argparse,hashlib,io,json
import unicodedata
from datetime import datetime,date
from decimal import Decimal
from pathlib import Path
import openpyxl
from openpyxl.utils.cell import range_boundaries

class InputRow(dict):
    def get(self,key,default=None):
        value=super().get(key,default)
        if key in self.formula_fields:
            raise ValueError(f'Formula in input field {key} at {self.source_range}; review the literal source before migration.')
        return value
    def __getitem__(self,key):
        if key not in self:raise KeyError(key)
        return self.get(key)

def uncertain_date(basis):
    normalized=''.join(c for c in unicodedata.normalize('NFD',str(basis or '').lower()) if unicodedata.category(c)!='Mn')
    return any(word in normalized for word in ('uoc','tam','random','du kien','chua xac nhan','chua ro'))

def run(source,output):
    raw=source.read_bytes()
    wb=openpyxl.load_workbook(io.BytesIO(raw),data_only=False)
    def table(name):
        for ws in wb:
            if name not in ws.tables:continue
            left,top,right,bottom=range_boundaries(ws.tables[name].ref)
            headers=[ws.cell(top,c).value for c in range(left,right+1)]
            records=[]
            for r in range(top+1,bottom+1):
                if ws.cell(r,left).value in (None,''):
                    core_fields={
                        't_suppliers':['Tên NCC'],
                        't_products':['Tên sản phẩm','NCC mặc định','Giá mua chuẩn'],
                        't_purchases':['SKU','Supplier_ID','Ngày thực nhận','SL thực nhận','Đơn giá mua'],
                        't_cash':['Số tiền thực tế','Ngày giờ','Reference_ID','Tài khoản chính'],
                    }.get(name,headers)
                    if any(ws.cell(r,c).data_type!='f' and ws.cell(r,c).value not in (None,'',0,False)
                           for h,c in zip(headers,range(left,right+1)) if h in core_fields):
                        raise ValueError(f'Missing source ID in {name} at {ws.title}!{r}; row contains input values and must be reviewed, not skipped.')
                    continue
                row=InputRow({h:ws.cell(r,c).value for h,c in zip(headers,range(left,right+1))})
                row['_source_cell']=f'{ws.title}!{openpyxl.utils.get_column_letter(left)}{r}:{openpyxl.utils.get_column_letter(right)}{r}'
                row.source_range=dict.get(row,'_source_cell')
                row.formula_fields={h for h,c in zip(headers,range(left,right+1)) if ws.cell(r,c).data_type=='f'}
                records.append(row)
            return records
        raise ValueError(f'Missing required Excel Table {name}')
    def numeric(value,label,minimum=0):
        n=Decimal(str(value or 0))
        if not n.is_finite() or n!=n.to_integral_value() or n<minimum or n>9_000_000_000_000:raise ValueError(f'{label}: invalid integer {value}')
        return int(n)
    def day(value):
        if value in (None,''):return None
        if isinstance(value,(datetime,date)):return value.strftime('%Y-%m-%d')
        raise ValueError(f'Invalid literal date: {value}')
    category_map={'PACKAGING':'packaging','SOFTWARE':'software','FLIVE':'software','RENT':'rent','UTILITIES':'utilities','SHIPPING':'shipping',
        'MARKETING':'marketing','PAYROLL':'payroll','OTHER':'other_expense','COGS_PURCHASE':'legacy_purchase_payment',
        'LEGACY_PURCHASE_PAYMENT':'legacy_purchase_payment','CAPITAL':'capital','WITHDRAWAL':'owner_withdrawal',
        'LEGACY_COD_RECEIPT':'legacy_cod','COD_RECEIPT':'legacy_cod','CUSTOMER_RECEIPT':'customer_receipt'}
    payload={'format_version':1,'source_id':hashlib.sha256(raw).hexdigest(),'source_name':source.name,
        'captured_at':datetime.now().isoformat(),'opening_date_assertion':'2026-08-18',
        'suppliers':[],'products':[],'purchases':[],'cash':[]}
    for r in table('t_suppliers'):
        payload['suppliers'].append({'code':r['Supplier_ID'],'name':r['Tên NCC'],'note':r.get('Ghi chú') or ''})
    for r in table('t_products'):
        payload['products'].append({'code':r['SKU'],'name':r['Tên sản phẩm'],'supplier_code':r.get('NCC mặc định'),
            'unit_cost':numeric(r.get('Giá mua chuẩn'),'Giá mua SKU'),'provisional':True,'note':r.get('Ghi chú') or ''})
    for r in table('t_purchases'):
        extras=sum(numeric(r.get(k),k) for k in ['Ship đầu vào','Gia công','Chi phí wash','Tag và label','Chi phí khác','Chi phí phân bổ'])
        discount=numeric(r.get('Chiết khấu dòng'),'Chiết khấu')
        if discount:raise ValueError('V1 needs an explicit discount field before migrating discounted purchases. No silent cost adjustment.')
        basis=r.get('Căn cứ ngày nhận') or ''
        payload['purchases'].append({'legacy_id':r['PO_Line_ID'],'supplier_code':r.get('Supplier_ID'),'product_code':r.get('SKU'),
            'received_date':day(r.get('Ngày thực nhận')),'date_estimated':uncertain_date(basis),'qty':numeric(r.get('SL thực nhận'),'Số lượng',1),
            'unit_cost':numeric(r.get('Đơn giá mua'),'Giá mua'),'additional_cost':extras,'notes':r.get('Ghi chú') or '',
            'source_status':r.get('Xác nhận nhận hàng'),'provenance':{'source_range':r['_source_cell'],'po_id':r.get('PO_ID'),
                'source_sku':r.get('SKU trong nguồn'),'date_basis':basis,'original_source':r.get('Tham chiếu nguồn gốc')}})
    for r in table('t_cash'):
        direction={'Thu':'in','Chi':'out'}.get(r.get('Thu Chi Transfer'))
        if not direction:raise ValueError('V1 does not migrate transfer records; needs a balanced transfer workflow.')
        basis=r.get('Căn cứ ngày và thanh toán') or ''
        category=category_map.get(r.get('Mã nhóm thu chi'))
        if category=='other_expense' and direction=='in':category='other_receipt'
        payload['cash'].append({'legacy_id':r['Cash_ID'],'direction':direction,'account_code':r.get('Tài khoản chính'),
            'transaction_date':day(r.get('Ngày giờ')),'date_estimated':uncertain_date(basis),'category':category,
            'amount':numeric(r.get('Số tiền thực tế'),'Số tiền',1),'description':r.get('Ghi chú') or r.get('Reference_ID') or r['Cash_ID'],
            'notes':r.get('Ghi chú') or '', 'source_status':r.get('Trạng thái ghi sổ'),
            'provenance':{'source_range':r['_source_cell'],'reference_id':r.get('Reference_ID'),'source_category':r.get('Mã nhóm thu chi'),
                'period_proposed':day(r.get('Kỳ lịch sử đề nghị')),'date_basis':basis,'original_source':r.get('Tham chiếu nguồn gốc')}})
    wb.close()
    output.parent.mkdir(parents=True,exist_ok=True)
    if output.exists():raise FileExistsError(f'Output exists; choose a new filename to keep the previous export: {output}')
    output.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':str(output),'source_sha256':payload['source_id'],'suppliers':len(payload['suppliers']),'products':len(payload['products']),
        'purchases':len(payload['purchases']),'cash':len(payload['cash']),'purchase_quantity':sum(r['qty'] for r in payload['purchases']),
        'purchase_value':sum(r['qty']*r['unit_cost']+r['additional_cost'] for r in payload['purchases'])},ensure_ascii=False,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,required=True);parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args();run(args.source.resolve(),args.output.resolve())
