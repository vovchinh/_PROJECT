"""Regression tests for source integrity; synthetic files only, never user Excel."""
import contextlib,io,json,tempfile,unittest,sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent))
from datetime import datetime
import openpyxl
from openpyxl.worksheet.table import Table
from export_excel import run,uncertain_date

class ExportIntegrity(unittest.TestCase):
    def make_source(self,folder,missing=None,formula=False):
        wb=openpyxl.Workbook();wb.remove(wb.active)
        tables={
            't_suppliers':(['Supplier_ID','Tên NCC'],[['SUP-A','Nhà cung cấp giả']]),
            't_products':(['SKU','Tên sản phẩm','NCC mặc định','Giá mua chuẩn'],[['SKU-A','Sản phẩm giả','SUP-A',45000]]),
            't_purchases':(['PO_Line_ID','Supplier_ID','SKU','SL thực nhận','Đơn giá mua','Ngày thực nhận','Căn cứ ngày nhận'],[[None if missing=='purchase' else 'P-1','SUP-A','SKU-A',2,'=45000' if formula else 45000,datetime(2026,8,18),'Ước theo tuần']]),
            't_cash':(['Cash_ID','Thu Chi Transfer','Số tiền thực tế','Ngày giờ','Căn cứ ngày và thanh toán','Mã nhóm thu chi'],[[None if missing=='cash' else 'C-1','Chi',10000,datetime(2026,8,18),'Ngày tạm','PACKAGING']]),
        }
        for name,(headers,rows) in tables.items():
            ws=wb.create_sheet(name);ws.append(headers)
            for row in rows:ws.append(row)
            # Template rows with no input are allowed.
            ws.append([None]*len(headers));ws.add_table(Table(displayName=name,ref=f'A1:{openpyxl.utils.get_column_letter(len(headers))}3'))
        path=Path(folder)/'synthetic.xlsx';wb.save(path);wb.close();return path
    def test_missing_purchase_id_is_not_silently_dropped(self):
        with tempfile.TemporaryDirectory() as folder:
            source=self.make_source(folder,missing='purchase');out=Path(folder)/'out.json'
            with self.assertRaisesRegex(ValueError,'Missing source ID'):run(source,out)
            self.assertFalse(out.exists())
    def test_missing_cash_id_is_not_silently_dropped(self):
        with tempfile.TemporaryDirectory() as folder:
            source=self.make_source(folder,missing='cash');out=Path(folder)/'out.json'
            with self.assertRaisesRegex(ValueError,'Missing source ID'):run(source,out)
            self.assertFalse(out.exists())
    def test_input_formula_requires_review(self):
        with tempfile.TemporaryDirectory() as folder:
            source=self.make_source(folder,formula=True)
            with self.assertRaisesRegex(ValueError,'Formula in input'):run(source,Path(folder)/'out.json')
    def test_literal_values_source_bytes_and_uncertainty_are_preserved(self):
        with tempfile.TemporaryDirectory() as folder:
            source=self.make_source(folder);before=source.read_bytes();out=Path(folder)/'out.json'
            with contextlib.redirect_stdout(io.StringIO()):run(source,out)
            payload=json.loads(out.read_text(encoding='utf-8'))
            self.assertEqual(source.read_bytes(),before);self.assertEqual(len(payload['purchases']),1);self.assertEqual(len(payload['cash']),1)
            self.assertEqual(payload['purchases'][0]['qty']*payload['purchases'][0]['unit_cost'],90000)
            self.assertTrue(payload['purchases'][0]['date_estimated']);self.assertTrue(payload['cash'][0]['date_estimated'])
            with self.assertRaises(FileExistsError):run(source,out)
    def test_shared_uncertainty_parser(self):
        for basis in ['Ước theo tuần','Ngày tạm','random','Chưa xác nhận','Dự kiến','ước tính']:
            self.assertTrue(uncertain_date(basis),basis)
        self.assertFalse(uncertain_date('Đã kiểm tra phiếu gốc'))

if __name__=='__main__':unittest.main(verbosity=2)
